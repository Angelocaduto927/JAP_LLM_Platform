import re
import os
from openpyxl import Workbook, load_workbook
from jap_question_generator_v3_1_qwen_based import clear_folder

def main():
    input_file = "docs/Generated_paper/shatin"
    output_file = "docs/Generated_paper/question_num_significance_test/paper"
    split_number = int(input("Enter the number of questions per file (e.g., 10): "))
    if os.path.exists(output_file):
        clear_folder(output_file)
    
    excel_files = os.listdir(input_file)
    for excel in excel_files:
        if excel.endswith(".xlsx"):
            process_file = os.path.join(input_file, excel)
            workbook1 = load_workbook(process_file)
            sheet1 = workbook1.active
            for index, row in enumerate(sheet1.iter_rows(values_only=True)):
                if index == 0:
                    headers = row
                else:
                    if (index-1) % split_number == 0:
                        match = re.match(r"^(.*)_(\d+)_(.*)\.xlsx", excel)
                        if match:
                            new_file_name = f"{match.group(1)}_{match.group(2)}_{(index-1)//split_number+1}_{match.group(3)}.xlsx"
                            workbook2 = Workbook()
                            sheet2 = workbook2.active
                            sheet2.append(headers)
                            output_path = os.path.join(output_file, new_file_name)
                            workbook2.save(output_path)
                    workbook2 = load_workbook(output_path)
                    sheet2 = workbook2.active
                    sheet2.append(row)
                    workbook2.save(output_path)
    return split_number

if __name__ == '__main__':
    main()