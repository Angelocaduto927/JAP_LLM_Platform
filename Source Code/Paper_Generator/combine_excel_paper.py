import re
import os
import shutil
import time
from openpyxl import Workbook, load_workbook
from jap_excel_processor import parse_excel_to_text

def combine_excel_files(selected_model, is_thinking_model, dir, split_number):
    excel = os.listdir(dir)
    record_dict = {}  # key: (prefix, group_key:int, suffix) -> list[(sort_key:int, filename)]
    # 兼容三种输入：3_4投票最终件 / 3_3最终件 / 旧的revised_revised
    pattern_voting_final = re.compile(r"^(.*)_(\d+)_(\d+)_(.*)_voting_group_\d+_revised\.xlsx$")
    pattern_single_final = re.compile(r"^(.*)_(\d+)_(\d+)_(.*)_revised\.xlsx$")
    pattern_legacy_rr = re.compile(r"^(.*)_(\d+)_(\d+)_(.*)_revised_revised\.xlsx$")

    for excel_file in excel:
        if not excel_file.endswith(".xlsx") or excel_file.startswith("~$"):
            continue
        if "_iteration_" in excel_file:
            # 跳过中间产物
            continue

        m = pattern_voting_final.match(excel_file)
        if m:
            prefix, group_key, sort_key, suffix = m.group(1), int(m.group(2)), int(m.group(3)), m.group(4)
            key = (prefix, group_key, suffix)
            record_dict.setdefault(key, []).append((sort_key, excel_file))
            continue

        m = pattern_single_final.match(excel_file)
        if m:
            prefix, group_key, sort_key, suffix = m.group(1), int(m.group(2)), int(m.group(3)), m.group(4)
            key = (prefix, group_key, suffix)
            record_dict.setdefault(key, []).append((sort_key, excel_file))
            continue

        m = pattern_legacy_rr.match(excel_file)
        if m:
            prefix, group_key, sort_key, suffix = m.group(1), int(m.group(2)), int(m.group(3)), m.group(4)
            key = (prefix, group_key, suffix)
            record_dict.setdefault(key, []).append((sort_key, excel_file))
            continue

    if not record_dict:
        print(f"Warning: No matching files found in {dir}")
        print("Expecting per-split outputs like:")
        print("  - *_voting_group_X_revised.xlsx (from v3_4)")
        print("  - *_revised.xlsx (from v3_3)")
        return

    tempt_path = os.path.join(dir, "tempt")
    if not os.path.exists(tempt_path):
        os.mkdir(tempt_path)

    for key, items in record_dict.items():
        # 按第三段（切分序号）排序
        items.sort(key=lambda x: x[0])
        full_qa = []
        # 逐个合并
        for sort_key, excel_file in items:
            new_dir = os.path.join(dir, excel_file)

            if full_qa == []:
                workbook = load_workbook(new_dir)
                worksheet = workbook.active
                header = [[cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]]
                full_qa.extend(header)

            qa = parse_excel_to_text(new_dir)
            for question in qa:
                # 规范列：将题号整理为 Qx 形式，并合并选项与答案列
                match = re.match(r'(Q\d+)', question[0])
                if match:
                    rest = question[0].replace(match.group(1), '').strip()
                    question[1] = rest + " " + question[1]
                    question[0] = match.group(1)
                    question[2] = "1. "+question[2]+" 2. "+question[3]+" 3. "+question[4]+" 4. "+question[5]
                    question[3] = question[6]
                    question[:] = question[:4]  # 截断到4列
            full_qa.extend(qa)

        prefix, group_key, suffix = key
        output_name_revised = f"{prefix}_{group_key}_{suffix}_revised.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        for row in full_qa:
            sheet.append(row)
        workbook.save(os.path.join(tempt_path, output_name_revised))
        workbook.close()
        time.sleep(0.05)
        print(f"Created combined file: {output_name_revised}")

    # 清理原目录xlsx，再移动新文件
    for excel_file in list(os.listdir(dir)):
        file_path = os.path.join(dir, excel_file)
        if os.path.isfile(file_path) and excel_file.endswith(".xlsx"):
            os.remove(file_path)

    for excel_file in os.listdir(tempt_path):
        src = os.path.join(tempt_path, excel_file)
        dst = os.path.join(dir, excel_file)
        shutil.move(src, dst)

    os.rmdir(tempt_path)

def main(split_number=10, model_index_input = None):
    print("====================================")
    print("如果是单独运行此脚本，试卷默认是按照10道题一份进行拆分的，如果需要更改分题数量，请用pipeline方式运行split_excel_paper.py，再运行此脚本// 或者可以在此脚本中修改split_number的默认值")
    print("====================================")
    
    # 定义模型配置（加入 voting_group_1..4 占位）
    model_config = {
        "qwen3-max": {"is_thinking": False},
        "qwen3-max-preview": {"is_thinking": False},
        "qwen-plus": {"is_thinking": True}, 
        "qwen3-vl-235b-a22b-instruct": {"is_thinking": False},
        "qwen-flash": {"is_thinking": True},
        "qwen3-30b-a3b-instruct-2507": {"is_thinking": False},
        "qwen-mt-plus": {"is_thinking": False},
        "qwen3-30b-a3b": {"is_thinking": True},
        "qwen3-32b": {"is_thinking": True},
        "qwen3-vl-235b-a22b-thinking": {"is_thinking": True},
        "qwen3-235b-a22b-thinking-2507":{"is_thinking": True},
        "qwen3-next-80b-a3b-thinking":{"is_thinking": True},
        "qwen3-next-80b-a3b-instruct":{"is_thinking": False},
        "qwen3-235b-a22b-instruct-2507":{"is_thinking": False},
        "qwen3-235b-a22b":{"is_thinking": False},
        "qwen3-30b-a3b-thinking-2507":{"is_thinking": True},
        # voting 组（不参与思考开关，仅用于定位目录）
        "voting_group_1": {"is_thinking": False},
        "voting_group_2": {"is_thinking": False},
        "voting_group_3": {"is_thinking": False},
        "voting_group_4": {"is_thinking": False},
    }
    
    if model_index_input is not None:
        selected_model = model_index_input
        if selected_model in model_config:
            is_thinking_model = model_config[selected_model]["is_thinking"]
        elif selected_model.startswith("voting_group_"):
            # 兜底：未在字典里仍允许，通过目录名使用
            is_thinking_model = False
        else:
            print(f"Unknown model: {selected_model}. Defaulting is_thinking_model=False")
            is_thinking_model = False
    else:
        available_models = list(model_config.keys())
        print("Available models:")
        for i, model in enumerate(available_models, 1):
            model_type = "Thinking" if model_config[model]["is_thinking"] else "Standard"
            print(f"{i}. {model} [{model_type}]")
        while True:
            try:
                choice = input(f"\nPlease select a model (1-{len(available_models)}): ").strip()
                model_index = int(choice) - 1
                if 0 <= model_index < len(available_models):
                    selected_model = available_models[model_index]
                    is_thinking_model = model_config[selected_model]["is_thinking"]
                    break
                else:
                    print("Invalid choice. Please try again.")
            except ValueError:
                print("Please enter a valid number.")
    
    print(f"\nSelected model: {selected_model}")
    print(f"Model type: {'Thinking mode' if is_thinking_model else 'Standard mode'}")
    
    # 输出目录：单模型与 voting_group_* 都是按名称分目录
    dir = f"docs/revised_shatin/{selected_model}"
    if not os.path.exists(dir):
        print(f"can't find directory: {dir}, run the model for revision first!")
        return
    
    combine_excel_files(selected_model, is_thinking_model, dir, split_number)
    print(f"\nAll Excel files have been processed using model: {selected_model}")

if __name__ == "__main__":
    main()