import pandas as pd
import os
import re
import shutil
from openpyxl import load_workbook, Workbook
from jap_excel_processor import parse_questions, parse_excel_to_text

def main():
    input_path = "docs/paper_with_feedback/4.11shatin"
    temp_path = "docs/paper_with_feedback/temp"
    final_path = "docs/paper_with_feedback/question_index"

    if not os.path.exists(final_path):
        os.makedirs(final_path)
    if not os.path.exists(temp_path):
        os.makedirs(temp_path)

    for file in os.listdir(input_path):
        if file.endswith('.xlsx') and not file.startswith('~$'):
            df = pd.read_excel(os.path.join(input_path, file))
            df = df[df["Suggestions"].isin(['Drop', 'Minor changes','OK'])]
            list_record = ["0"]*len(df)
            bad_questions_df = df[df['Suggestions'].isin(['Drop', 'Minor changes'])]

            if os.path.exists(os.path.join(temp_path, file)):
                os.remove(os.path.join(temp_path, file))
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(bad_questions_df.columns.tolist())
            for index, row in bad_questions_df.iterrows():
                sheet.append(row.tolist())
            workbook.save(os.path.join(temp_path, file))
            
            qa = parse_excel_to_text(os.path.join(temp_path, file))
            bad_questions_index = []
            for item in qa:
                bad_questions_index.append(item[0])
                match = re.match(r'Q(\d+): もんだい(\d+)', item[0])
                if match:
                    question_num = match.group(1)
                    problem_num = match.group(2)
                    list_record[(int(problem_num)-1)*10+int(question_num)-1] = "1"
                else:
                    print(f"Unexpected question format: {item[0]}")
            
            string_record = "".join(list_record) 
            
            if os.path.exists(os.path.join(final_path, file)):
                os.remove(os.path.join(final_path, file))
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Bad Questions Index", "String_record"])
            sheet.append([", ".join(bad_questions_index), string_record])
            workbook.save(os.path.join(final_path, file))
            print(f"{file} processed successfully.")

    shutil.rmtree(temp_path)
    
if __name__ == "__main__":
    main()